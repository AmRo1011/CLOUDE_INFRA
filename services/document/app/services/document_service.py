"""
Document service
"""
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func
from sqlalchemy.orm import selectinload
from fastapi import UploadFile
from typing import Optional, List, Tuple
import logging
import uuid

from app.models.document import Document, DocumentContent
from app.s3_client import s3_client
from app.kafka_client import kafka_producer
from app.config import get_settings

logger = logging.getLogger(__name__)
settings = get_settings()


class DocumentService:
    def __init__(self, db: AsyncSession):
        self.db = db
    
    async def upload_document(
        self, 
        user_id: str, 
        file: UploadFile, 
        title: str, 
        description: Optional[str] = None
    ) -> Document:
        """Upload and process a document"""
        # Validate file
        ext = file.filename.split(".")[-1].lower() if file.filename else ""
        if ext not in settings.allowed_extensions_list:
            raise ValueError(f"File type not allowed. Allowed: {settings.ALLOWED_EXTENSIONS}")
        
        file_size = 0
        content = await file.read()
        file_size = len(content)
        
        if file_size > settings.MAX_UPLOAD_SIZE_MB * 1024 * 1024:
            raise ValueError(f"File too large. Max: {settings.MAX_UPLOAD_SIZE_MB}MB")
        
        # Generate S3 key
        doc_id = str(uuid.uuid4())
        s3_key = f"documents/{user_id}/{doc_id}/{file.filename}"
        
        # Upload to S3
        from io import BytesIO
        await s3_client.upload_file(BytesIO(content), s3_key, file.content_type or "application/octet-stream")
        
        # Create document record
        document = Document(
            id=doc_id,
            user_id=user_id,
            title=title,
            description=description,
            file_name=file.filename,
            file_type=file.content_type,
            file_size=file_size,
            s3_key=s3_key,
            status="processing"
        )
        self.db.add(document)
        await self.db.flush()
        
        # Extract text and store
        extracted_text = await self._extract_text(content, ext)
        if extracted_text:
            doc_content = DocumentContent(document_id=doc_id, content=extracted_text)
            self.db.add(doc_content)
            document.status = "processed"
        else:
            document.status = "uploaded"
        
        await self.db.flush()
        await self.db.refresh(document)
        
        # Publish event
        await kafka_producer.send_event(
            topic="document.uploaded",
            payload={
                "document_id": doc_id,
                "user_id": user_id,
                "title": title,
                "file_name": file.filename,
                "s3_key": s3_key
            }
        )
        
        if document.status == "processed":
            await kafka_producer.send_event(
                topic="document.processed",
                payload={
                    "document_id": doc_id,
                    "user_id": user_id,
                    "title": title,
                    "content_length": len(extracted_text) if extracted_text else 0
                }
            )
        
        return document
    
    async def _extract_text(self, content: bytes, ext: str) -> Optional[str]:
        """Extract text from document"""
        try:
            if ext == "pdf":
                from PyPDF2 import PdfReader
                from io import BytesIO
                reader = PdfReader(BytesIO(content))
                text = ""
                for page in reader.pages:
                    text += page.extract_text() or ""
                return text.strip()
            elif ext == "docx":
                from docx import Document as DocxDocument
                from io import BytesIO
                doc = DocxDocument(BytesIO(content))
                return "\n".join([p.text for p in doc.paragraphs])
            elif ext in ["txt", "md"]:
                return content.decode("utf-8", errors="ignore")
        except Exception as e:
            logger.error(f"Text extraction failed: {e}")
        return None
    
    async def get_document(self, document_id: str, user_id: str) -> Optional[Document]:
        """Get document by ID"""
        result = await self.db.execute(
            select(Document)
            .options(selectinload(Document.content))
            .where(Document.id == document_id, Document.user_id == user_id)
        )
        return result.scalar_one_or_none()
    
    async def list_documents(
        self, 
        user_id: str, 
        page: int = 1, 
        limit: int = 20, 
        status: Optional[str] = None
    ) -> Tuple[List[Document], int]:
        """List user's documents"""
        query = select(Document).where(Document.user_id == user_id)
        count_query = select(func.count(Document.id)).where(Document.user_id == user_id)
        
        if status:
            query = query.where(Document.status == status)
            count_query = count_query.where(Document.status == status)
        
        total_result = await self.db.execute(count_query)
        total = total_result.scalar() or 0
        
        offset = (page - 1) * limit
        query = query.offset(offset).limit(limit).order_by(Document.created_at.desc())
        
        result = await self.db.execute(query)
        return list(result.scalars().all()), total
    
    async def get_download_url(self, s3_key: str) -> Optional[str]:
        """Get presigned download URL"""
        return await s3_client.get_presigned_url(s3_key)
    
    async def delete_document(self, document_id: str, user_id: str) -> bool:
        """Delete document"""
        document = await self.get_document(document_id, user_id)
        if not document:
            return False
        
        await s3_client.delete_file(document.s3_key)
        await self.db.delete(document)
        await self.db.flush()
        return True

